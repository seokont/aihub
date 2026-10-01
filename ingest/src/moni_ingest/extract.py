"""Text extraction — one function per format, and no silent empties.

Every extractor returns the document's plain text and raises :class:`ExtractionError` when it
cannot produce any. That matters more than it looks: an extractor that returns `""` for a
scanned PDF would be *ingested* as a source with zero chunks, and the operator would only
discover it when a question about that document found nothing. Failing at ingest time names
the file instead.

Supported: PDF (pypdf, with pdfplumber as a fallback), DOCX, and Markdown/text.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Final

import structlog

log = structlog.get_logger(__name__)

#: Extensions this package will ingest. Anything else is refused rather than guessed at —
#: OCR for images and spreadsheets are Phase 3 concerns.
SUPPORTED_SUFFIXES: Final[frozenset[str]] = frozenset({".pdf", ".docx", ".md", ".markdown", ".txt"})

#: Collapse runs of blank lines. PDF extraction in particular emits a blank line per layout
#: gap, which inflates the text without adding meaning and wastes tokens on every query.
_BLANK_RUNS: Final = re.compile(r"\n{3,}")
#: Trailing spaces before a newline, which PDF/DOCX extraction leaves behind.
_TRAILING_SPACE: Final = re.compile(r"[ \t]+\n")


class ExtractionError(RuntimeError):
    """The file could not be read as text."""

    def __init__(self, path: Path, reason: str) -> None:
        super().__init__(f"{path.name}: {reason}")
        self.path = path
        self.reason = reason


@dataclass(frozen=True, slots=True)
class Extracted:
    """Plain text plus the facts worth recording about where it came from."""

    text: str
    #: Which reader produced the text. Recorded so a later quality question ("why is this
    #: document's text poor?") can be answered without re-running the extraction.
    extractor: str
    pages: int | None = None


def normalise(text: str) -> str:
    """Tidy extracted text without changing its words.

    Deliberately conservative: no de-hyphenation and no whitespace re-flowing, because both
    would alter text that a user may later quote back. Only whitespace noise is removed.
    """
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _TRAILING_SPACE.sub("\n", text)
    text = _BLANK_RUNS.sub("\n\n", text)
    return text.strip()


def _extract_pdf(path: Path) -> Extracted:
    """PDF text via pypdf, falling back to pdfplumber.

    The fallback is not belt-and-braces for its own sake: pypdf returns an empty string for
    some producer/encoding combinations where pdfplumber succeeds, and that specific failure
    is invisible until retrieval finds nothing.
    """
    from pypdf import PdfReader

    pages: list[str] = []
    with path.open("rb") as handle:
        reader = PdfReader(handle)
        for page in reader.pages:
            pages.append(page.extract_text() or "")
    text = normalise("\n\n".join(pages))
    if text:
        return Extracted(text=text, extractor="pypdf", pages=len(pages))

    log.info("pdf_pypdf_empty_falling_back", file=path.name, pages=len(pages))
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        pages = [page.extract_text() or "" for page in pdf.pages]
    text = normalise("\n\n".join(pages))
    if not text:
        msg = "no extractable text (a scanned or image-only PDF needs OCR, which is Phase 3)"
        raise ExtractionError(path, msg)
    return Extracted(text=text, extractor="pdfplumber", pages=len(pages))


def _extract_docx(path: Path) -> Extracted:
    """DOCX text: paragraphs plus table cells.

    Table contents are included because client documents carry their figures in tables, and
    skipping them would drop exactly the numbers people ask about.
    """
    import docx

    document = docx.Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = normalise("\n".join(part for part in parts if part.strip()))
    if not text:
        raise ExtractionError(path, "the document has no text paragraphs or tables")
    return Extracted(text=text, extractor="python-docx")


def _extract_plain(path: Path) -> Extracted:
    """Markdown/plain text, read as UTF-8 with a tolerant fallback."""
    try:
        text = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = path.read_text(encoding="utf-8", errors="replace")
        log.warning("text_decoded_with_replacements", file=path.name)
    text = normalise(text)
    if not text:
        raise ExtractionError(path, "the file is empty")
    return Extracted(text=text, extractor="plain")


def extract(path: Path) -> Extracted:
    """Extract text from ``path``, dispatching on its suffix."""
    if not path.is_file():
        raise ExtractionError(path, "not a file")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise ExtractionError(path, f"unsupported type {suffix!r} (supported: {supported})")
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    return _extract_plain(path)


def discover(target: Path) -> list[Path]:
    """Expand a file or directory into the ingestible files beneath it.

    Sorted, so a directory ingest is deterministic and two runs produce the same order —
    which is what lets "re-run → zero new rows" be checked without sorting noise.
    """
    if target.is_file():
        return [target]
    if not target.is_dir():
        raise ExtractionError(target, "not a file or directory")
    return sorted(
        path
        for path in target.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
    )


__all__ = [
    "SUPPORTED_SUFFIXES",
    "Extracted",
    "ExtractionError",
    "discover",
    "extract",
    "normalise",
]
